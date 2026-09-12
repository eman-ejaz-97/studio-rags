#!/usr/bin/env python3
"""Generate the Studio Rags Requirements and Deliverables Mapping (v2.0)."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

HEADERS = [
    "User story",
    "Functional requirement",
    "Non-functional requirement",
    "Acceptance criteria",
    "Related deliverable",
]
WIDTHS = [Cm(5.4), Cm(6.2), Cm(5.4), Cm(6.2), Cm(3.9)]

# (section title, [rows]) — each row is the 5 columns above.
CONTENT = [
    ("A. Browsing and booking", [
        (
            "As a customer, I want to view the workshops on offer so that I can "
            "choose one that interests me.",
            "FR-C1, FR-C2 — The system shall display all published offers with "
            "name, image, description, duration, price and upcoming dates, "
            "grouped by type: one-off workshop, four-week Master Class, online "
            "class, and kids class.",
            "NFR-2, NFR-3 — The page shall be responsive and load within 3 "
            "seconds on a normal mobile connection.",
            "All 9 offers are listed with correct information: Indonesian "
            "Batik, Japanese Shibori and Woodblock Printing as one-off "
            "workshops; the three matching Master Classes; Batik and Shibori "
            "online; and Kids & Parents Painting.",
            "Workshop listing page, workshop cards, UI prototype.",
        ),
        (
            "As a customer, I want to see available dates and remaining places "
            "so that I can choose a session that suits me.",
            "FR-C3 — The system shall display upcoming session dates, times and "
            "the number of places still available for each offer.",
            "NFR-1, NFR-9 — The calendar shall be clear, keyboard accessible and "
            "usable on phone, tablet and desktop.",
            "Available sessions can be selected. Full, cancelled and past "
            "sessions cannot be selected.",
            "Workshop calendar, session selection page.",
        ),
        (
            "As a customer, I want to book a workshop so that I can reserve my "
            "place.",
            "FR-C6 — The system shall allow a customer to select a session, "
            "enter participant and contact details, choose the number of places "
            "and submit a booking.",
            "NFR-1 — The booking process shall be completed in a small number of "
            "clearly explained steps, and shall be completable one-handed on a "
            "phone.",
            "A booking is created only when all required details are entered and "
            "enough places remain. The places chosen are removed from "
            "availability immediately.",
            "Booking form, booking module, booking database tables.",
        ),
        (
            "As a customer, I want to pay securely online so that my booking is "
            "confirmed straight away.",
            "FR-C7, FR-P1, FR-P2, FR-P3, FR-P5 — The system shall take full "
            "payment through Stripe before confirming a booking, and shall "
            "record the transaction reference and payment status. No deposits "
            "or part payments.",
            "NFR-4 — All payment traffic shall use HTTPS. Card details shall "
            "never be stored in the Studio Rags database; they are handled "
            "entirely by Stripe.",
            "The booking changes to confirmed only after Stripe reports a "
            "successful payment. The customer then sees an on-screen "
            "confirmation and receives a confirmation email. A failed or "
            "abandoned payment never produces a confirmed booking.",
            "Checkout page, Stripe integration, booking confirmation page.",
        ),
        (
            "As a customer, I want to book an online class so that I can take "
            "part without travelling to the studio.",
            "FR-C15 — The system shall offer Batik and Shibori as online "
            "classes, and shall provide joining details to the customer after "
            "payment.",
            "NFR-4 — Joining details shall be visible only to the customer who "
            "booked that session.",
            "An online booking is confirmed after payment and the joining "
            "details appear in the customer's account and confirmation email.",
            "Online session booking flow, joining details email.",
        ),
        (
            "As a parent, I want to book the Kids & Parents Painting class so "
            "that I can attend with my child.",
            "FR-C6 — The system shall support the Kids & Parents Painting class "
            "for children aged 7 to 15, and shall capture the child's details at "
            "booking.",
            "NFR-7 — Children's details shall be collected only where necessary "
            "and handled in line with the privacy notice.",
            "A parent can book the class, enter the child's details, and pay. "
            "The admin sees both the parent and child on the booking.",
            "Kids class booking flow, admin booking view.",
        ),
    ]),
    ("B. Master Class (new in this version)", [
        (
            "As a customer, I want to buy a four-week Master Class so that I can "
            "learn a craft in depth.",
            "FR-C16 — The system shall sell a four-week Master Class for Batik, "
            "Shibori and Woodblock Printing, paid in full in advance, valid for "
            "six months from purchase.",
            "NFR-4 — Payment shall follow the same secure process as a normal "
            "booking.",
            "After successful payment the customer holds an enrolment recording "
            "four sessions and an expiry date six months ahead. A confirmation "
            "email is sent.",
            "Master Class purchase page, enrolment records.",
        ),
        (
            "As a Master Class student, I want to book my four sessions so that "
            "I can attend at times that suit me.",
            "FR-C17 — The system shall allow an enrolled student to book "
            "sessions against the published schedule, using one of their four "
            "sessions instead of paying again.",
            "NFR-5 — An enrolled student shall occupy a place in exactly the "
            "same way as a paying customer, so a class can never be "
            "overbooked.",
            "Booking a session reduces the student's remaining sessions by one "
            "and reduces the class's available places by one. A student cannot "
            "book more than four, and cannot book after the expiry date.",
            "Session redemption flow, enrolment records.",
        ),
        (
            "As a Master Class student, I want to see how many sessions I have "
            "left so that I do not lose them.",
            "FR-C18, FR-C19 — The system shall show remaining sessions and the "
            "expiry date in the customer's account, and shall email a reminder "
            "before the enrolment expires.",
            "NFR-1 — The remaining sessions and expiry date shall be visible "
            "without the customer having to search for them.",
            "The account page shows sessions used, sessions remaining and the "
            "expiry date. A reminder email is sent before expiry.",
            "Customer account dashboard, expiry reminder email.",
        ),
    ]),
    ("C. Customer account and communication", [
        (
            "As a customer, I want to register and log in so that I can view and "
            "manage my bookings.",
            "FR-C4 — The system shall allow a customer to create an account, log "
            "in, log out and view their booking history.",
            "NFR-4 — Passwords shall be securely hashed. Login errors shall be "
            "clear without revealing whether an account exists.",
            "A registered customer can log in with valid details and see only "
            "their own bookings. Invalid credentials are rejected.",
            "Registration page, login page, customer account dashboard.",
        ),
        (
            "As a customer, I want to reset a forgotten password so that I can "
            "get back into my account.",
            "FR-C5 — The system shall send a password reset link to the "
            "registered email address on request.",
            "NFR-4 — Reset links shall expire after a short period and shall "
            "work once only.",
            "A customer receives a reset link, sets a new password, and can log "
            "in with it. The old password no longer works.",
            "Password reset pages, reset email.",
        ),
        (
            "As a customer, I want to view my booking history so that I can keep "
            "track of what I have booked.",
            "FR-C10 — The system shall show the customer their past and upcoming "
            "bookings, with date, status and amount paid.",
            "NFR-4 — A customer shall never be able to see another customer's "
            "bookings.",
            "The account page lists the customer's bookings correctly, showing "
            "the amount they actually paid at the time of booking.",
            "Customer account dashboard.",
        ),
        (
            "As a customer, I want to receive a confirmation and a reminder so "
            "that I do not forget my workshop.",
            "FR-C8, FR-C9, FR-N1, FR-N2 — The system shall send an automatic "
            "confirmation email on booking, and automatic reminder emails before "
            "the session.",
            "NFR-5 — Emails shall be sent reliably and every send shall be "
            "recorded so it can be checked later.",
            "A confirmation email arrives after payment. A reminder email "
            "arrives before the session. Both contain the correct date, time and "
            "location or joining details.",
            "Email templates, email sending service, email log.",
        ),
        (
            "As a customer, I want to leave a review so that I can share my "
            "experience with others.",
            "FR-C11 — The system shall allow a customer who has attended a "
            "session to leave a rating and a comment, and shall display approved "
            "reviews on the workshop page.",
            "NFR-7 — Reviews shall be moderated before they appear publicly.",
            "Only a customer with an attended booking can review that workshop, "
            "and only once. Reviews appear publicly only after the admin "
            "approves them.",
            "Review form, review moderation screen, workshop page reviews.",
        ),
        (
            "As a customer, I want to buy and redeem a gift card so that I can "
            "give a workshop as a present.",
            "FR-C12, FR-P4 — The system shall allow a customer to choose a gift "
            "card value, enter recipient details, pay, and later redeem a valid "
            "code at checkout.",
            "NFR-4 — Gift card codes shall be unique, securely generated and "
            "protected against reuse.",
            "A digital gift card is issued only after successful payment. A "
            "valid unused code reduces the amount payable. Invalid, spent or "
            "expired codes are rejected.",
            "Gift card purchase page, digital voucher, redemption module.",
        ),
        (
            "As a customer, I want to contact Studio Rags so that I can ask a "
            "question before booking.",
            "FR-C13 — The system shall provide a contact page with a form that "
            "sends an enquiry to the admin.",
            "NFR-4 — The contact form shall be protected against automated "
            "spam.",
            "A submitted enquiry reaches the admin and is recorded in the "
            "system. The customer sees a confirmation on screen.",
            "Contact page, enquiry records.",
        ),
    ]),
    ("D. Administration", [
        (
            "As the administrator, I want to see today's and this week's classes "
            "so that I can prepare.",
            "FR-A2, FR-A10 — The system shall provide a dashboard showing "
            "upcoming classes, a calendar, customer names, contact details, "
            "booking status and payment status.",
            "NFR-1, NFR-6 — The dashboard shall be simple enough to use without "
            "training or developer help.",
            "On opening the dashboard the admin can see upcoming classes and who "
            "is attending, without searching.",
            "Admin dashboard.",
        ),
        (
            "As the administrator, I want to create and update workshops and "
            "dates so that customers see accurate information.",
            "FR-A3, FR-A4 — The system shall allow the admin to create, edit, "
            "publish, cancel and remove workshops and individual dates.",
            "NFR-4, NFR-6 — Only authorised administrators shall have access. "
            "Changes shall be saved reliably.",
            "An admin can add or change a workshop or a date, and the change "
            "appears correctly on the public website.",
            "Admin dashboard, workshop management module.",
        ),
        (
            "As the administrator, I want to set the price and the number of "
            "places myself so that I do not need a developer.",
            "FR-A11 — The system shall allow the admin to set a default price "
            "and number of places for each workshop, and to change either for an "
            "individual date.",
            "NFR-6 — Changing a price shall never change what an existing "
            "customer has already paid.",
            "A price or place count set by the admin appears immediately on the "
            "website. Existing bookings continue to show the amount originally "
            "paid.",
            "Workshop and session management screens.",
        ),
        (
            "As the administrator, I want to be sure a class can never be "
            "overbooked so that I do not have to turn people away.",
            "FR-A5 — The system shall prevent bookings beyond the number of "
            "places available, including when several customers book at the same "
            "moment.",
            "NFR-5 — The system shall behave correctly under the higher booking "
            "volumes of the spring and summer season.",
            "In a concurrency test, 50 simultaneous attempts on a class with 10 "
            "places produce exactly 10 confirmed bookings and no more.",
            "Booking module, seat locking logic, load test report.",
        ),
        (
            "As the administrator, I want to view and manage bookings so that I "
            "can prepare for upcoming classes.",
            "FR-A6, FR-A7, FR-A8 — The system shall display booking details, "
            "customer name, contact details, number of participants, session, "
            "payment status and booking status, with filtering.",
            "NFR-3, NFR-4 — Customer data shall be accessible only to authorised "
            "users. Searching and filtering shall respond within 3 seconds.",
            "The admin can filter bookings by workshop, date and status, and can "
            "update or cancel a booking where permitted.",
            "Booking management dashboard, booking report.",
        ),
        (
            "As the administrator, I want to see who holds a Master Class "
            "enrolment so that I can manage the courses.",
            "FR-A12 — The system shall list Master Class enrolments showing the "
            "student, sessions used, sessions remaining and expiry date.",
            "NFR-6 — The admin shall be able to see and adjust an enrolment "
            "without developer help.",
            "The admin can find any enrolment, see its remaining sessions and "
            "expiry date, and record an adjustment where needed.",
            "Enrolment management screen.",
        ),
        (
            "As the administrator, I want to be notified of new bookings so that "
            "I always know what is coming up.",
            "FR-A9, FR-N3, FR-N4 — The system shall email the admin when a "
            "booking is made, and before an upcoming class.",
            "NFR-5 — Notification emails shall be sent reliably and recorded.",
            "The admin receives an email for each new booking, and a reminder "
            "before each upcoming class.",
            "Admin notification emails, email log.",
        ),
    ]),
    ("E. Marketing, privacy and quality", [
        (
            "As the business owner, I want to see which bookings came from my "
            "adverts so that I know whether my advertising works.",
            "FR-M1 — The system shall be built so that a Meta or Google tracking "
            "pixel can be added, and shall report booking and payment events to "
            "it.",
            "NFR-7 — Tracking shall run only where the visitor has given "
            "consent.",
            "With a pixel installed, a completed booking is reported to the ad "
            "platform. With consent declined, no tracking occurs.",
            "Tracking setup, consent banner, pixel configuration guide.",
        ),
        (
            "As the business owner, I want to collect customer email addresses "
            "so that I can tell people about new workshops.",
            "FR-M2, FR-M3 — The system shall collect marketing consent "
            "separately from booking, provide a privacy and cookie notice, and "
            "allow the admin to export the consented list.",
            "NFR-7 — Consent shall be explicit and recorded with a date. "
            "Personal data shall be handled in line with Australian privacy "
            "expectations.",
            "A customer can book without agreeing to marketing. Only customers "
            "who opted in appear in the exported list, each with a recorded "
            "consent date.",
            "Consent capture, privacy notice, customer export.",
        ),
        (
            "As a customer with limited eyesight or dexterity, I want the site "
            "to be easy to use so that I can book without difficulty.",
            "FR-C14 — The site shall be designed for phones first and shall work "
            "across common browsers and devices.",
            "NFR-9 — The site shall meet WCAG 2.2 Level AA: readable text sizes, "
            "sufficient colour contrast, visible focus indicators and full "
            "keyboard operation.",
            "An accessibility audit records no Level AA failures on the booking "
            "journey. The whole booking can be completed using a keyboard "
            "alone.",
            "Design system, accessibility audit report.",
        ),
    ]),
]


def shade(cell, hex_colour):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def write_cell(cell, text, *, size=8.5, bold=False, colour=INK, space_after=0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour


doc = Document()

section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Cm(29.7), Cm(21.0)
section.left_margin = section.right_margin = Cm(1.2)
section.top_margin = section.bottom_margin = Cm(1.2)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

title = doc.add_paragraph()
run = title.add_run("Studio Rags Capstone Project | Requirements and Deliverables Mapping")
run.font.size = Pt(16)
run.font.bold = True
run.font.color.rgb = ACCENT
title.paragraph_format.space_after = Pt(2)

meta = doc.add_paragraph()
run = meta.add_run("Version 2.0  ·  4 August 2026  ·  Aligned to SRS v1.0")
run.font.size = Pt(9)
run.font.color.rgb = MUTED
meta.paragraph_format.space_after = Pt(8)

intro = doc.add_paragraph()
run = intro.add_run(
    "This table links each user story to its functional requirement, "
    "non-functional requirement, acceptance criteria and related project "
    "deliverable. Requirement codes match the Software Requirements "
    "Specification version 1.0."
)
run.font.size = Pt(9.5)
intro.paragraph_format.space_after = Pt(6)

changes = doc.add_paragraph()
run = changes.add_run("What changed since version 1.0 of this table")
run.font.size = Pt(10)
run.font.bold = True
run.font.color.rgb = ACCENT
changes.paragraph_format.space_after = Pt(2)

for line in [
    "The four-week Master Class has been added. It is paid in full in advance "
    "and taken within six months, so it is described in its own section (B).",
    "The offer list is now the final nine: three one-off workshops, three "
    "Master Classes, two online classes and the kids class. Tie-dye and Chunri "
    "have been removed.",
    "Prices and class sizes are now set by the client herself, and are recorded "
    "as a requirement rather than as information to be gathered.",
    "Stories that were missing have been added: password reset, confirmation "
    "and reminder emails, reviews, the contact page, admin notifications, "
    "overbooking prevention, privacy and consent, and accessibility.",
    "Stripe has been confirmed as the payment provider.",
]:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(line)
    r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(1)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

hdr = table.rows[0]
repeat_header(hdr)
for i, name in enumerate(HEADERS):
    hdr.cells[i].width = WIDTHS[i]
    write_cell(hdr.cells[i], name, size=9.5, bold=True, colour=RGBColor(0xFF, 0xFF, 0xFF))
    shade(hdr.cells[i], "1F3A5F")

for section_title, rows in CONTENT:
    band = table.add_row()
    for i in range(5):
        band.cells[i].width = WIDTHS[i]
        shade(band.cells[i], "E8EDF4")
    merged = band.cells[0].merge(band.cells[4])
    write_cell(merged, section_title, size=10, bold=True, colour=ACCENT)

    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            row.cells[i].width = WIDTHS[i]
            write_cell(row.cells[i], value)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
foot = doc.add_paragraph()
run = foot.add_run(
    "Still to be confirmed with the client: how the Master Class runs in the "
    "studio (section B), the cancellation and refund policy, how online class "
    "materials are provided, and how Kids & Parents Painting tickets are "
    "priced. These are listed in section 7 of the SRS v1.0."
)
run.font.size = Pt(9)
run.font.color.rgb = MUTED

out = "/Users/eman/Projects/studio_rags/docs/client/Studio_Rags_Requirements_and_Deliverables_v2.docx"
doc.save(out)
print("saved:", out)
print("stories:", sum(len(r) for _, r in CONTENT))
